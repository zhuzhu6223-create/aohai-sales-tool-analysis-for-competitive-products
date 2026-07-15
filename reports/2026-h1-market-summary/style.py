"""Shared style tokens + helper primitives for the deck. Mirrors Hallmark's
locked-token discipline: every color/font used in build.py must come from here."""
import math
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- tokens -----------------------------------------------------------
PAPER = RGBColor(0xF7, 0xF5, 0xF0)
PAPER_2 = RGBColor(0xFC, 0xFB, 0xF8)
PAPER_3 = RGBColor(0xEF, 0xEB, 0xE1)
BORDER = RGBColor(0xDF, 0xD9, 0xCB)
INK = RGBColor(0x1E, 0x2A, 0x38)
INK_2 = RGBColor(0x54, 0x60, 0x70)
INK_3 = RGBColor(0x87, 0x8F, 0x98)

PRIMARY = RGBColor(0x3A, 0x50, 0x6B)      # 深水蓝
PRIMARY_SOFT = RGBColor(0xDE, 0xE4, 0xEA)
PRIMARY_DEEP = RGBColor(0x27, 0x36, 0x49)

COMPLETE = RGBColor(0x5B, 0x9E, 0x9A)     # 柔青绿
COMPLETE_SOFT = RGBColor(0xE1, 0xED, 0xEB)

ASSET = RGBColor(0xA8, 0xB5, 0x9F)        # 鼠尾草绿
ASSET_SOFT = RGBColor(0xEC, 0xEF, 0xE7)

RISK = RGBColor(0xC9, 0x7B, 0x47)         # 暖橙 — cap ~8% area per slide
RISK_SOFT = RGBColor(0xF5, 0xE6, 0xD8)

DISPLAY_FONT = "微软雅黑"       # headings — bold weight carries the display role
BODY_FONT = "微软雅黑"          # body / data
MONO_FONT = "Consolas"          # numerals, page tags, version labels

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.56)
CONTENT_W = SLIDE_W - 2 * MARGIN

DECK_TITLE = "遨海科技市场部 · 2026 年半年工作总结"


def E(v):
    """Coerce any numeric (possibly float, e.g. from `/` division on Length
    ints) to a valid integer-EMU Length. python-pptx's low-level oxml setters
    do not always cast for you, and a float EMU value serializes as
    "2560320.0" — schema-invalid OOXML that PowerPoint/LibreOffice reject."""
    return Emu(int(round(v)))


def set_bg(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _no_shadow(shape):
    shape.shadow.inherit = False


def rect(slide, left, top, width, height, fill=None, line=None, line_w=0.75,
         shape=MSO_SHAPE.RECTANGLE, dash=None):
    left, top, width, height = E(left), E(top), E(width), E(height)
    shp = slide.shapes.add_shape(shape, left, top, width, height)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash:
            ln = shp.line._get_or_add_ln()
            d = ln.makeelement(qn('a:prstDash'), {'val': dash})
            ln.append(d)
    _no_shadow(shp)
    return shp


def oval(slide, cx, cy, r, fill=None, line=None, line_w=1.0):
    return rect(slide, cx - r, cy - r, r * 2, r * 2, fill=fill, line=line,
                line_w=line_w, shape=MSO_SHAPE.OVAL)


def connector_line(slide, x1, y1, x2, y2, color=BORDER, weight=1.0, dash=None):
    x1, y1, x2, y2 = E(x1), E(y1), E(x2), E(y2)
    ln = slide.shapes.add_connector(1, x1, y1, x2, y2)  # 1 = straight
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    if dash:
        el = ln.line._get_or_add_ln()
        d = el.makeelement(qn('a:prstDash'), {'val': dash})
        el.append(d)
    _no_shadow(ln)
    return ln


def rich(slide, left, top, width, height, paragraphs, anchor=MSO_ANCHOR.TOP,
         wrap=True):
    left, top, width, height = E(left), E(top), E(width), E(height)
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get('align', PP_ALIGN.LEFT)
        p.line_spacing = para.get('line_spacing', 1.18)
        p.space_after = Pt(para.get('space_after', 6))
        p.space_before = Pt(para.get('space_before', 0))
        for run in para['runs']:
            r = p.add_run()
            r.text = run['text']
            r.font.size = Pt(run.get('size', 13))
            r.font.bold = run.get('bold', False)
            r.font.italic = False  # typography purity — no italic headers/runs
            r.font.name = run.get('font', BODY_FONT)
            r.font.color.rgb = run.get('color', INK)
    return tb


def text(slide, left, top, width, height, s, size=13, color=INK, bold=False,
         font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.18, space_after=6):
    paras = [{'runs': [{'text': line, 'size': size, 'color': color,
                         'bold': bold, 'font': font}],
              'align': align, 'line_spacing': line_spacing,
              'space_after': space_after}
             for line in s.split('\n')]
    return rich(slide, left, top, width, height, paras, anchor=anchor)


def bullets(slide, left, top, width, height, items, size=12.5, color=INK,
            marker_color=PRIMARY, space_after=7, line_spacing=1.2,
            bold_lead=None):
    """items: list of str, or list of (lead_bold_text, rest_text)."""
    paras = []
    for it in items:
        runs = [{'text': '— ', 'size': size, 'color': marker_color,
                 'bold': True, 'font': BODY_FONT}]
        if isinstance(it, tuple):
            lead, rest = it
            runs.append({'text': lead, 'size': size, 'color': color,
                         'bold': True, 'font': BODY_FONT})
            if rest:
                runs.append({'text': rest, 'size': size, 'color': color,
                             'bold': False, 'font': BODY_FONT})
        else:
            runs.append({'text': it, 'size': size, 'color': color,
                         'bold': False, 'font': BODY_FONT})
        paras.append({'runs': runs, 'space_after': space_after,
                      'line_spacing': line_spacing})
    return rich(slide, left, top, width, height, paras)


def chip(slide, left, top, width, height, label, fill=PRIMARY_SOFT,
         text_color=PRIMARY, size=10.5, bold=True, line=None, align=PP_ALIGN.CENTER):
    rect(slide, left, top, width, height, fill=fill, line=line, line_w=0.75,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(slide, left, top, width, height, label, size=size, color=text_color,
         bold=bold, align=align, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05,
         space_after=0)


def badge_num(slide, cx, cy, r, n, fill=PRIMARY, text_color=PAPER, size=13):
    oval(slide, cx, cy, r, fill=fill)
    text(slide, cx - r, cy - r, r * 2, r * 2, str(n), size=size,
         color=text_color, bold=True, font=MONO_FONT, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, space_after=0)


# ---- page chrome --------------------------------------------------------

def header(slide, page_no, kicker, title, conclusion=None, title_size=25,
           title_color=INK, two_line_title=None):
    # page tag, top-left, mono
    text(slide, MARGIN, Inches(0.32), Inches(2.4), Inches(0.3),
         f"{page_no:02d} / 12", size=10.5, color=INK_3, font=MONO_FONT,
         bold=False, space_after=0)
    # wordmark, top-right
    wm_w, wm_h = Inches(1.5), Inches(0.32)
    rect(slide, SLIDE_W - MARGIN - wm_w, Inches(0.30), wm_w, wm_h,
         fill=None, line=BORDER, line_w=0.75)
    text(slide, SLIDE_W - MARGIN - wm_w, Inches(0.30), wm_w, wm_h, "遨海科技",
         size=10.5, color=INK_2, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    # kicker
    ky = Inches(0.66)
    if kicker:
        text(slide, MARGIN, ky, CONTENT_W, Inches(0.26), kicker, size=11,
             color=PRIMARY, bold=True, font=MONO_FONT, space_after=0)
        ty = Inches(0.94)
    else:
        ty = Inches(0.68)
    # title
    if two_line_title:
        paras = [{'runs': [{'text': two_line_title[0], 'size': title_size,
                             'bold': True, 'color': title_color,
                             'font': DISPLAY_FONT}], 'space_after': 2,
                  'line_spacing': 1.08}]
        paras.append({'runs': [{'text': two_line_title[1], 'size': title_size,
                                 'bold': True, 'color': title_color,
                                 'font': DISPLAY_FONT}], 'space_after': 4,
                       'line_spacing': 1.08})
        th = Inches(0.95)
        rich(slide, MARGIN, ty, CONTENT_W, th, paras)
        ty = ty + th
    else:
        th = Inches(0.55)
        text(slide, MARGIN, ty, CONTENT_W, th, title, size=title_size,
             bold=True, color=title_color, font=DISPLAY_FONT,
             line_spacing=1.12, space_after=0)
        ty = ty + th + Inches(0.06)
    cur_y = ty
    if conclusion:
        bar_h = Inches(0.5) if len(conclusion) < 62 else Inches(0.62)
        rect(slide, MARGIN, cur_y, CONTENT_W, bar_h, fill=PRIMARY_SOFT,
             shape=MSO_SHAPE.RECTANGLE)
        rect(slide, MARGIN, cur_y, Inches(0.05), bar_h, fill=PRIMARY)
        text(slide, MARGIN + Inches(0.2), cur_y, CONTENT_W - Inches(0.35),
             bar_h, conclusion, size=12.5, color=PRIMARY_DEEP, bold=False,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15, space_after=0)
        cur_y = cur_y + bar_h + Inches(0.16)
    else:
        cur_y = cur_y + Inches(0.1)
    return cur_y


def footer(slide, page_no, note=None):
    y = SLIDE_H - Inches(0.36)
    connector_line(slide, MARGIN, y, SLIDE_W - MARGIN, y, color=BORDER, weight=0.75)
    text(slide, MARGIN, y + Inches(0.05), Inches(8), Inches(0.26), DECK_TITLE,
         size=9, color=INK_3, font=MONO_FONT, space_after=0)
    if note:
        text(slide, Inches(4.2), y + Inches(0.05), Inches(6), Inches(0.26),
             note, size=9, color=INK_3, align=PP_ALIGN.CENTER, space_after=0)
    text(slide, SLIDE_W - MARGIN - Inches(1.2), y + Inches(0.05), Inches(1.2),
         Inches(0.26), f"{page_no:02d}", size=9, color=INK_3, font=MONO_FONT,
         align=PP_ALIGN.RIGHT, space_after=0)


def bottom_conclusion(slide, s, top=None, fill=PRIMARY, text_color=PAPER, size=14):
    h = Inches(0.62)
    if top is None:
        top = SLIDE_H - Inches(0.36) - Inches(0.14) - h
    rect(slide, MARGIN, top, CONTENT_W, h, fill=fill)
    text(slide, MARGIN + Inches(0.28), top, CONTENT_W - Inches(0.56), h, s,
         size=size, color=text_color, bold=True, anchor=MSO_ANCHOR.MIDDLE,
         line_spacing=1.15, space_after=0)
    return top


def card(slide, left, top, width, height, fill=PAPER_2, line=BORDER, line_w=0.75):
    return rect(slide, left, top, width, height, fill=fill, line=line,
                line_w=line_w, shape=MSO_SHAPE.ROUNDED_RECTANGLE)


def polar(cx, cy, r, deg):
    rad = math.radians(deg)
    return Emu(int(cx + r * math.cos(rad))), Emu(int(cy + r * math.sin(rad)))
